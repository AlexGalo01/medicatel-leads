from __future__ import annotations

import csv
from pathlib import Path
from typing import Any
from uuid import UUID

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill


def export_leads_to_csv(job_id: UUID, leads: list[dict[str, Any]], export_dir_path: str) -> str:
    export_dir = Path(export_dir_path)
    export_dir.mkdir(parents=True, exist_ok=True)

    export_path = export_dir / f"leads_{job_id}.csv"
    columns = [
        "full_name",
        "specialty",
        "country",
        "city",
        "score",
        "score_reasoning",
        "email",
        "whatsapp",
        "linkedin_url",
    ]

    with export_path.open(mode="w", encoding="utf-8", newline="") as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=columns)
        writer.writeheader()
        for lead in leads:
            writer.writerow({column: lead.get(column, "") for column in columns})

    return str(export_path)


def export_leads_to_xlsx(job_id: UUID, leads: list[dict[str, Any]], export_dir_path: str) -> str:
    """Export leads to Excel format with formatting."""
    export_dir = Path(export_dir_path)
    export_dir.mkdir(parents=True, exist_ok=True)

    export_path = export_dir / f"leads_{job_id}.xlsx"

    columns = [
        ("full_name", "Nombre"),
        ("specialty", "Especialidad"),
        ("country", "País"),
        ("city", "Ciudad"),
        ("score", "Puntuación"),
        ("score_reasoning", "Razonamiento"),
        ("email", "Email"),
        ("whatsapp", "WhatsApp"),
        ("phone", "Teléfono"),
        ("linkedin_url", "LinkedIn"),
        ("address", "Dirección"),
        ("schedule_text", "Horario"),
        ("primary_source_url", "URL fuente"),
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leads"

    # Header row — bold white text on primary color
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="6366F1", end_color="6366F1", fill_type="solid")
    for col_idx, (_, header_label) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header_label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Data rows
    for row_idx, lead in enumerate(leads, start=2):
        for col_idx, (field, _) in enumerate(columns, start=1):
            value = lead.get(field) or ""
            ws.cell(row=row_idx, column=col_idx, value=value)

    # Auto column width
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    # Freeze header row
    ws.freeze_panes = "A2"

    wb.save(export_path)
    return str(export_path)
